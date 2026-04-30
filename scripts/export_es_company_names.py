#!/usr/bin/env python3
"""Generate an Excel training file augmented with real ES English company names."""

from __future__ import annotations

import argparse
import os
import random
import re
import sys
from pathlib import Path
from typing import Iterable, Optional
from urllib import parse as urllib_parse

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from export_es_broker_entities import _hit_values, _json_request, _resolve_path


ASCII_ALPHA_RE = re.compile(r"[A-Za-z]")


def _dedupe_english_names(values: Iterable[str]) -> list[str]:
    names, seen = [], set()
    for value in values:
        name = re.sub(r"\s+", " ", str(value)).strip()
        if not name or not ASCII_ALPHA_RE.search(name) or name in seen:
            continue
        names.append(name)
        seen.add(name)
    return names


def fetch_company_names(args: argparse.Namespace) -> list[str]:
    base_url = f"http://{args.host}:{args.port}"
    index = urllib_parse.quote(args.index, safe="")
    search_url = f"{base_url}/{index}/_search?scroll={urllib_parse.quote(args.scroll)}"
    payload = {
        "size": args.batch_size,
        "_source": [args.field],
        "query": {"exists": {"field": args.field}},
    }
    response = _json_request(
        search_url,
        user=args.user,
        password=args.password,
        payload=payload,
        timeout=args.timeout,
    )

    names = []
    scroll_id: Optional[str] = response.get("_scroll_id")
    try:
        while True:
            hits = (response.get("hits") or {}).get("hits") or []
            if not hits:
                break
            for hit in hits:
                names.extend(_hit_values(hit, args.field))

            if args.max_names and len(names) >= args.max_names:
                break
            if not scroll_id:
                break
            response = _json_request(
                f"{base_url}/_search/scroll",
                user=args.user,
                password=args.password,
                payload={"scroll": args.scroll, "scroll_id": scroll_id},
                timeout=args.timeout,
            )
            scroll_id = response.get("_scroll_id", scroll_id)
    finally:
        if scroll_id:
            try:
                _json_request(
                    f"{base_url}/_search/scroll",
                    user=args.user,
                    password=args.password,
                    payload={"scroll_id": [scroll_id]},
                    method="DELETE",
                    timeout=args.timeout,
                )
            except RuntimeError:
                pass

    names = _dedupe_english_names(names)
    return names[: args.max_names] if args.max_names else names


def _header_index(headers: list[str], name: str) -> int:
    try:
        return headers.index(name)
    except ValueError as exc:
        raise ValueError(f"Missing required column: {name}") from exc


def _company_entity_lines(value) -> list[tuple[int, list[str]]]:
    if value is None:
        return []
    lines = []
    for line_index, raw_line in enumerate(str(value).splitlines()):
        parts = [part.strip() for part in raw_line.split("|")]
        if len(parts) >= 2 and parts[0] == "机构-公司" and parts[1]:
            lines.append((line_index, parts))
    return lines


def _replace_company_entity_cell(value, line_index: int, replacement: str) -> str:
    raw_lines = str(value or "").splitlines()
    parts = [part.strip() for part in raw_lines[line_index].split("|")]
    parts[1] = replacement
    if len(parts) >= 3:
        parts[2] = replacement
    raw_lines[line_index] = "|".join(parts)
    return "\n".join(raw_lines)


def augment_workbook(args: argparse.Namespace, company_names: list[str]) -> int:
    input_path = _resolve_path(args.input_file)
    output_path = _resolve_path(args.output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(input_path)
    worksheet = workbook.active
    header_values = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
    ]
    query_idx = _header_index(header_values, args.query_column)
    entities_idx = _header_index(header_values, args.company_entities_column)
    rng = random.Random(args.seed)

    source_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    rng.shuffle(source_rows)
    augmented = 0

    for values in source_rows:
        if args.max_examples is not None and augmented >= args.max_examples:
            break
        if rng.random() >= args.prob:
            continue

        row_values = list(values)
        while len(row_values) < len(header_values):
            row_values.append(None)
        query = str(row_values[query_idx] or "").strip()
        entity_cell = row_values[entities_idx]
        candidates = [
            (line_index, parts)
            for line_index, parts in _company_entity_lines(entity_cell)
            if parts[1] in query
        ]
        if not query or not candidates:
            continue

        line_index, parts = rng.choice(candidates)
        source_mention = parts[1]
        replacement = rng.choice(company_names)
        row_values[query_idx] = query.replace(source_mention, replacement, 1)
        row_values[entities_idx] = _replace_company_entity_cell(entity_cell, line_index, replacement)
        worksheet.append(row_values)
        augmented += 1

    workbook.save(output_path)
    return augmented


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Fetch base_company.en_company_name from Elasticsearch and append "
            "English-company augmented rows to an Excel training file."
        )
    )
    parser.add_argument("--host", default=os.getenv("ELASTICSEARCH_SEARCH_URL", "10.100.0.100"))
    parser.add_argument(
        "--port",
        type=int,
        default=int(os.getenv("ELASTICSEARCH_SEARCH_PORT", "9200")),
    )
    parser.add_argument("--user", default=os.getenv("ELASTICSEARCH_SEARCH_USER", "elastic"))
    parser.add_argument("--password", default=os.getenv("ELASTICSEARCH_SEARCH_PASSWORD", "6YCuCbNUf2ap"))
    parser.add_argument("--index", default="base_company")
    parser.add_argument("--field", default="en_company_name")
    parser.add_argument(
        "--input-file",
        default="data/test_data_0413.xlsx",
        help="Source annotated Excel file, relative to project root unless absolute.",
    )
    parser.add_argument(
        "--output-file",
        default="data/test_data_0413_en_augmented.xlsx",
        help="Output augmented Excel file, relative to project root unless absolute.",
    )
    parser.add_argument("--query-column", default="Query")
    parser.add_argument("--company-entities-column", default="company_entities")
    parser.add_argument("--prob", type=float, default=1.0)
    parser.add_argument("--max-examples", type=int, default=5000)
    parser.add_argument("--max-names", type=int, default=None)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--batch-size", type=int, default=1000)
    parser.add_argument("--scroll", default="2m")
    parser.add_argument("--timeout", type=float, default=30.0)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        company_names = fetch_company_names(args)
    except RuntimeError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    if not company_names:
        print(f"no English company names found in {args.index}.{args.field}", file=sys.stderr)
        return 1

    try:
        augmented = augment_workbook(args, company_names)
    except (OSError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 1

    print(
        f"loaded {len(company_names)} English company names; "
        f"appended {augmented} augmented rows to {_resolve_path(args.output_file)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
