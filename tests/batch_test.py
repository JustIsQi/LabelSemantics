import argparse
import json
import re
import sys
import time
from collections import defaultdict
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DEPENDENCY_ERROR = None

try:
    import torch
    from transformers import AutoTokenizer

    from label_semantics.labels import build_tag_maps, load_label_descriptions
    from label_semantics.metrics import get_entities
    from label_semantics.model import LabelSemanticsNER
except ModuleNotFoundError as exc:
    torch = None
    AutoTokenizer = None
    build_tag_maps = None
    load_label_descriptions = None
    get_entities = None
    LabelSemanticsNER = None
    DEPENDENCY_ERROR = exc

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as exc:
    Workbook = None
    Alignment = None
    Font = None
    PatternFill = None
    get_column_letter = None
    if DEPENDENCY_ERROR is None:
        DEPENDENCY_ERROR = exc


def parse_args():
    parser = argparse.ArgumentParser(description="Batch-test entity extraction with a self-contained model directory.")
    parser.add_argument("--input-file", default="data/eval.jsonl")
    parser.add_argument(
        "--model-path",
        default="outputs/best_model",
        help="Path to a self-contained model directory (e.g. produced by scripts/convert_checkpoint_to_model_dir.py).",
    )
    parser.add_argument("--label-file", default="data/excel_ner_data/labels.json")
    parser.add_argument("--output-file", default="tests/batch_test_predictions.xlsx")
    parser.add_argument("--max-length", type=int, default=128)
    parser.add_argument("--device", default=None)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--warmup", type=int, default=5)
    return parser.parse_args()


def load_jsonl(path, limit=None):
    rows = []
    with Path(path).open("r", encoding="utf-8") as input_file:
        for line_number, line in enumerate(input_file, start=1):
            if limit is not None and len(rows) >= limit:
                break
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSON on line {line_number}: {exc}") from exc
    return rows


# Characters that should never appear at the start/end of an entity span. These are
# typical separators between enumerated companies/brokers in the training data
# ("中信证券、中金公司"), plus generic whitespace and ASCII/CJK punctuation.
ENTITY_BOUNDARY_STRIP_CHARS = set(
    " \t\u3000、，,。.；;：:!！?？/\\()()[]【】《》<>\"'`~“”‘’\r\n"
)


def trim_entity_span(query, start, end):
    """Trim leading/trailing punctuation or whitespace from a predicted span.

    Returns (new_start, new_end). If everything is trimmed away, returns a span where
    ``new_start > new_end`` so the caller can drop it.
    """
    while start <= end and query[start] in ENTITY_BOUNDARY_STRIP_CHARS:
        start += 1
    while end >= start and query[end] in ENTITY_BOUNDARY_STRIP_CHARS:
        end -= 1
    return start, end


def encode_query(query, tokenizer, max_length, device):
    chars = list(query)
    inputs = tokenizer(
        chars,
        is_split_into_words=True,
        max_length=max_length,
        padding="max_length",
        truncation=True,
        return_tensors="pt",
    )
    word_ids = inputs.word_ids(batch_index=0)
    model_inputs = {
        "input_ids": inputs["input_ids"].to(device),
        "attention_mask": inputs["attention_mask"].to(device),
        "token_type_ids": inputs.get("token_type_ids", torch.zeros_like(inputs["input_ids"])).to(device),
    }
    return model_inputs, word_ids


def predict_entities(query, model, tokenizer, id2tag, label_descriptions, max_length, device):
    model_inputs, word_ids = encode_query(query, tokenizer, max_length, device)
    with torch.no_grad():
        _, pred_ids = model(**model_inputs, use_label_cache=True)

    tags = ["O"] * len(query)
    previous_word_id = None
    for pred_id, word_id in zip(pred_ids[0].detach().cpu().tolist(), word_ids):
        if word_id is None or word_id >= len(tags):
            previous_word_id = word_id
            continue
        # Align with training: only the FIRST sub-token of each word carries the supervised
        # label. Non-first sub-tokens are unsupervised (IGNORE_INDEX) at training time and
        # produce noisy predictions, so we ignore them at inference too.
        if word_id == previous_word_id:
            continue
        tags[word_id] = id2tag[pred_id]
        previous_word_id = word_id

    entities = {label: [] for label in label_descriptions}
    for entity_type, start, end in get_entities(tags):
        if entity_type not in entities:
            continue
        start, end = trim_entity_span(query, start, end)
        if start > end:
            continue
        entities[entity_type].append(query[start:end + 1])
    return entities


def normalize_mentions(mentions):
    return sorted({str(mention).strip() for mention in mentions if str(mention).strip()})


TIME_PATTERNS = (
    r"截至\d{4}年\d{1,2}月\d{1,2}日",
    r"\d{4}年\d{1,2}月\d{1,2}日",
    r"\d{4}\s*[-—至到]\s*\d{4}年",
    r"\d{4}年(?:第[一二三四1234]季度|[一二三四1234]季度|Q[1-4]|H[12]|上半年|下半年|中报|年报|全年)",
    r"\d{4}Q[1-4]",
    r"\d{4}年",
    r"\d{1,2}月\d{1,2}日",
    r"(?:最近|近|过去|未来|前)\d*[一二三四五六七八九十0-9]*(?:个)?(?:交易日|工作日|日|天|周|月|季度|季|半年|年)",
    r"(?:近|过去)[一二三四五六七八九十0-9]+年",
    r"年初至今",
    r"(?:最近|最新|近期|当前|目前|今年|去年|本周|本月|本季度|本年|上周|上月|上季度|上年|上半年|下半年|上市以来)",
    r"[一二三四]季度",
    r"Q[1-4]",
)
TIME_RE = re.compile("|".join(f"(?:{pattern})" for pattern in TIME_PATTERNS))


def extract_time_mentions(query):
    mentions = []
    seen = set()
    for match in TIME_RE.finditer(query):
        mention = match.group(0).strip()
        if mention and mention not in seen:
            mentions.append(mention)
            seen.add(mention)
    return mentions


def normalize_entities(entities, label_descriptions):
    return {
        label: normalize_mentions(entities.get(label, []))
        for label in label_descriptions
    }


def gold_entities_from_row(row, label_descriptions):
    raw_entities = row.get("label", {}).get("entities", {}) or {}
    keys_by_label = {
        "C": ("companies",),
        "BROKER": ("institutions",),
        "IND": ("industries",),
        "CODE": ("codes", "code"),
        "TIME": ("times", "time"),
    }
    query = row.get("ori_query", "")
    entities = {}
    for label in label_descriptions:
        values = []
        for key in keys_by_label.get(label, (label.lower(),)):
            values.extend(raw_entities.get(key, []))
        if label == "TIME":
            values.extend(extract_time_mentions(query))
        entities[label] = values
    return normalize_entities(
        entities,
        label_descriptions,
    )


def sync_if_cuda(device):
    if device.type == "cuda":
        torch.cuda.synchronize(device)


def evaluate_match(pred_entities, gold_entities):
    """Per-label scoring: a label is 'hit' iff gold ⊆ pred (i.e. all gold mentions
    appear in the prediction). Returns (all_hit, per_label_info) where per_label_info
    is {label: {"hit": bool, "extras": [...], "missing": [...]}}.
    """
    per_label = {}
    all_hit = True
    for label in gold_entities:
        gold_set = set(gold_entities.get(label, []))
        pred_set = set(pred_entities.get(label, []))
        missing = sorted(gold_set - pred_set)
        extras = sorted(pred_set - gold_set)
        hit = not missing
        per_label[label] = {"hit": hit, "extras": extras, "missing": missing}
        if not hit:
            all_hit = False
    return all_hit, per_label


def _format_mentions(mentions):
    return "\n".join(mentions) if mentions else ""


def write_excel_report(output_path, detail_rows, label_descriptions, summary, input_file, model_path):
    """Write a human-friendly Excel workbook with two sheets: predictions detail + summary."""
    workbook = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    wrong_fill = PatternFill("solid", fgColor="FCE4E4")
    extras_fill = PatternFill("solid", fgColor="FFF2CC")
    missing_fill = PatternFill("solid", fgColor="F8CBAD")
    correct_font = Font(color="2E7D32", bold=True)
    wrong_font = Font(color="C62828", bold=True)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    center_align = Alignment(horizontal="center", vertical="center")

    labels_sorted = sorted(label_descriptions.keys())

    # ---------- Sheet 1: Predictions detail ----------
    detail = workbook.active
    detail.title = "明细"

    headers = ["#", "原始Query", "是否正确"]
    for label in labels_sorted:
        zh = label_descriptions.get(label, label)
        headers.extend([
            f"{label}/{zh} - 标注",
            f"{label}/{zh} - 预测",
            f"{label}/{zh} - 多余",
            f"{label}/{zh} - 遗漏",
        ])
    headers.append("耗时(ms)")

    detail.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = detail.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    column_widths = [6, 42, 10]
    for _ in labels_sorted:
        column_widths.extend([18, 18, 18, 18])
    column_widths.append(10)
    for col_idx, width in enumerate(column_widths, start=1):
        detail.column_dimensions[get_column_letter(col_idx)].width = width

    for row in detail_rows:
        excel_row = [
            row["index"],
            row["query"],
            "✓" if row["correct"] else "✗",
        ]
        for label in labels_sorted:
            gold_list = row["gold"].get(label, [])
            pred_list = row["pred"].get(label, [])
            info = row["per_label"].get(label, {"extras": [], "missing": []})
            excel_row.extend([
                _format_mentions(gold_list),
                _format_mentions(pred_list),
                _format_mentions(info.get("extras", [])),
                _format_mentions(info.get("missing", [])),
            ])
        excel_row.append(round(row["elapsed_ms"], 3))
        detail.append(excel_row)

        excel_row_idx = detail.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = detail.cell(row=excel_row_idx, column=col_idx)
            cell.alignment = wrap_align
            if not row["correct"]:
                cell.fill = wrong_fill

        correct_cell = detail.cell(row=excel_row_idx, column=3)
        correct_cell.alignment = center_align
        correct_cell.font = correct_font if row["correct"] else wrong_font

        for label_idx, label in enumerate(labels_sorted):
            base_col = 4 + label_idx * 4  # cols: gold, pred, extras, missing
            extras_cell = detail.cell(row=excel_row_idx, column=base_col + 2)
            missing_cell = detail.cell(row=excel_row_idx, column=base_col + 3)
            if extras_cell.value:
                extras_cell.fill = extras_fill
            if missing_cell.value:
                missing_cell.fill = missing_fill

    detail.freeze_panes = "C2"
    detail.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ---------- Sheet 2: Summary ----------
    summary_sheet = workbook.create_sheet("汇总")

    summary_sheet.append(["指标", "数值"])
    for col_idx in (1, 2):
        cell = summary_sheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    overview_rows = [
        ("输入文件", str(input_file)),
        ("模型路径", str(model_path)),
        ("样本总数", summary["total"]),
        ("全对样本数", summary["correct"]),
        ("整体准确率(全对)", f"{summary['accuracy']:.6f}"),
        ("平均耗时(ms)", f"{summary['avg_ms']:.3f}"),
        ("最大耗时(ms)", f"{summary['max_ms']:.3f}"),
        ("含多余预测的样本数", summary["rows_with_any_extras"]),
        ("多余实体总数", summary["total_extras"]),
    ]
    for key, value in overview_rows:
        summary_sheet.append([key, value])

    summary_sheet.append([])
    title_row_idx = summary_sheet.max_row + 1
    per_label_headers = [
        "标签",
        "中文",
        "样本数",
        "命中数(gold⊆pred)",
        "召回率",
        "多余样本数",
        "多余实体数",
        "遗漏样本数",
        "遗漏实体数",
    ]
    summary_sheet.append(per_label_headers)
    for col_idx in range(1, len(per_label_headers) + 1):
        cell = summary_sheet.cell(row=title_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for label in labels_sorted:
        ttl = summary["label_total"].get(label, 0)
        hit = summary["label_hits"].get(label, 0)
        recall = hit / ttl if ttl else 0.0
        summary_sheet.append([
            label,
            label_descriptions.get(label, ""),
            ttl,
            hit,
            f"{recall:.4f}",
            summary["label_extras_rows"].get(label, 0),
            summary["label_extras_count"].get(label, 0),
            summary["label_missing_rows"].get(label, 0),
            summary["label_missing_count"].get(label, 0),
        ])

    summary_widths = [12, 12, 10, 18, 12, 14, 14, 14, 14]
    for col_idx, width in enumerate(summary_widths, start=1):
        summary_sheet.column_dimensions[get_column_letter(col_idx)].width = width

    workbook.save(output_path)


def main():
    args = parse_args()
    if DEPENDENCY_ERROR is not None:
        print(f"Missing dependency: {DEPENDENCY_ERROR.name}", file=sys.stderr)
        print("Install dependencies with: pip install -r requirements.txt", file=sys.stderr)
        raise SystemExit(1)

    device = torch.device(args.device or ("cuda" if torch.cuda.is_available() else "cpu"))

    print(f"Using device: {device}", flush=True)
    print(f"Loading labels from {args.label_file}", flush=True)
    label_descriptions = load_label_descriptions(args.label_file)
    tag2id, id2tag = build_tag_maps(label_descriptions)

    print(f"Loading tokenizer from {args.model_path}", flush=True)
    tokenizer = AutoTokenizer.from_pretrained(args.model_path, use_fast=True)
    print(f"Loading model from {args.model_path}", flush=True)
    model = LabelSemanticsNER(args.model_path, tag2id, label_descriptions).to(device)
    model.eval()

    rows = load_jsonl(args.input_file, args.limit)
    print(f"Loaded {len(rows)} rows from {args.input_file}", flush=True)

    for row in rows[:args.warmup]:
        predict_entities(
            row.get("ori_query", ""),
            model,
            tokenizer,
            id2tag,
            label_descriptions,
            args.max_length,
            device,
        )
    sync_if_cuda(device)

    output_path = Path(args.output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    correct = 0
    total = 0
    total_ms = 0.0
    max_ms = 0.0
    label_hits = defaultdict(int)
    label_total = defaultdict(int)
    label_extras_rows = defaultdict(int)
    label_extras_count = defaultdict(int)
    label_missing_rows = defaultdict(int)
    label_missing_count = defaultdict(int)
    rows_with_any_extras = 0
    total_extras = 0

    detail_rows = []

    for index, row in enumerate(rows, start=1):
        query = row.get("ori_query", "")
        gold_entities = gold_entities_from_row(row, label_descriptions)

        sync_if_cuda(device)
        start = time.perf_counter()
        pred_entities = normalize_entities(
            predict_entities(query, model, tokenizer, id2tag, label_descriptions, args.max_length, device),
            label_descriptions,
        )
        sync_if_cuda(device)
        elapsed_ms = (time.perf_counter() - start) * 1000

        is_correct, per_label = evaluate_match(pred_entities, gold_entities)
        correct += int(is_correct)
        total += 1
        total_ms += elapsed_ms
        max_ms = max(max_ms, elapsed_ms)

        row_extras_total = 0
        for label, info in per_label.items():
            label_total[label] += 1
            if info["hit"]:
                label_hits[label] += 1
            if info["extras"]:
                label_extras_rows[label] += 1
                label_extras_count[label] += len(info["extras"])
                row_extras_total += len(info["extras"])
            if info["missing"]:
                label_missing_rows[label] += 1
                label_missing_count[label] += len(info["missing"])
        if row_extras_total:
            rows_with_any_extras += 1
            total_extras += row_extras_total

        detail_rows.append({
            "index": index,
            "query": query,
            "correct": is_correct,
            "gold": gold_entities,
            "pred": pred_entities,
            "per_label": per_label,
            "elapsed_ms": elapsed_ms,
        })

        if index % 100 == 0:
            accuracy = correct / total if total else 0.0
            avg_ms = total_ms / total if total else 0.0
            print(
                f"processed={index} accuracy={accuracy:.6f} avg_ms={avg_ms:.3f}",
                flush=True,
            )

    accuracy = correct / total if total else 0.0
    avg_ms = total_ms / total if total else 0.0

    summary = {
        "total": total,
        "correct": correct,
        "accuracy": accuracy,
        "avg_ms": avg_ms,
        "max_ms": max_ms,
        "rows_with_any_extras": rows_with_any_extras,
        "total_extras": total_extras,
        "label_total": dict(label_total),
        "label_hits": dict(label_hits),
        "label_extras_rows": dict(label_extras_rows),
        "label_extras_count": dict(label_extras_count),
        "label_missing_rows": dict(label_missing_rows),
        "label_missing_count": dict(label_missing_count),
    }

    write_excel_report(
        output_path,
        detail_rows,
        label_descriptions,
        summary,
        input_file=args.input_file,
        model_path=args.model_path,
    )

    print(f"Output: {output_path}", flush=True)
    print(
        f"total={total} correct={correct} accuracy={accuracy:.6f} "
        f"avg_ms={avg_ms:.3f} max_ms={max_ms:.3f}",
        flush=True,
    )
    print("Per-label recall (gold ⊆ pred) and extras:", flush=True)
    for label in sorted(label_total):
        ttl = label_total[label]
        hit = label_hits[label]
        recall = hit / ttl if ttl else 0.0
        ex_rows = label_extras_rows[label]
        ex_cnt = label_extras_count[label]
        ms_rows = label_missing_rows[label]
        ms_cnt = label_missing_count[label]
        print(
            f"  {label:>8s}: recall={recall:.4f} ({hit}/{ttl})  "
            f"extras_rows={ex_rows} extras_total={ex_cnt}  "
            f"missing_rows={ms_rows} missing_total={ms_cnt}",
            flush=True,
        )
    print(
        f"Rows with any extras: {rows_with_any_extras}  total extras mentions: {total_extras}",
        flush=True,
    )


if __name__ == "__main__":
    main()
